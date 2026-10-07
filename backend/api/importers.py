"""
استيراد قطع الغيار من Excel/CSV وتصديرها.

المحل الذي يبدأ النظام عنده آلاف الأصناف في دفتر أو ملف Excel؛ إدخالها
يدوياً عائق دخول حقيقي. الاستيراد يقبل عناوين أعمدة عربية أو إنجليزية،
ويعرض معاينة بالأخطاء قبل أي حفظ، ثم يطبّق الملف كاملاً في معاملة واحدة.

القواعد:
- رقم القطعة هو المفتاح: رقم جديد = قطعة جديدة برصيدها الافتتاحي (حركة
  موثّقة). رقم موجود = تحديث بياناتها الوصفية وسعر البيع فقط عند اختيار
  التحديث؛ لا يُكتب فوق متوسط التكلفة ولا الرصيد (لهما مسارات التوريد والجرد).
- الفئة غير الموجودة تُنشأ تلقائياً.
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from . import services
from .models import Category, SparePart
from .search import normalize_text

MAX_ROWS = 20000

# العمود → أسماء عناوينه المقبولة (تُقارن بعد توحيد الحروف).
COLUMNS = {
    'name': ['name', 'الاسم', 'اسم القطعة', 'القطعة', 'الصنف'],
    'part_number': ['part_number', 'part number', 'رقم القطعة', 'الرقم', 'الكود', 'كود'],
    'category': ['category', 'الفئة', 'القسم', 'المجموعة'],
    'purchase_price': ['purchase_price', 'cost', 'سعر الشراء', 'التكلفة'],
    'selling_price': ['selling_price', 'price', 'سعر البيع', 'السعر'],
    'quantity': ['quantity', 'stock', 'الكمية', 'الرصيد', 'العدد'],
    'min_stock_alert': ['min_stock_alert', 'حد التنبيه', 'الحد الادنى'],
    'shelf_location': ['shelf_location', 'shelf', 'الرف', 'موقع الرف'],
    'brand': ['brand', 'العلامه', 'الماركه', 'المصنع', 'الشركه'],
    'oem_number': ['oem', 'oem_number', 'الرقم الاصلي'],
    'quality_grade': ['quality', 'quality_grade', 'الجوده', 'درجه الجوده', 'النوع'],
    'aliases': ['aliases', 'اسماء اخري', 'الاسم الدارج', 'اسماء دارجه'],
    'barcode': ['barcode', 'الباركود'],
    'description': ['description', 'الوصف'],
}
HEADER_LOOKUP = {
    normalize_text(alias): column for column, aliases in COLUMNS.items() for alias in aliases
}

GRADE_LOOKUP = {
    normalize_text(label): value
    for value, labels in {
        SparePart.QualityGrade.ORIGINAL: ['original', 'أصلي', 'اصلي', 'وكالة'],
        SparePart.QualityGrade.COMMERCIAL: ['commercial', 'تجاري'],
        SparePart.QualityGrade.USED: ['used', 'مستعمل', 'تشليح'],
    }.items()
    for label in labels
}

# ترتيب أعمدة ملف التصدير والقالب (عناوين عربية).
EXPORT_COLUMNS = [
    ('part_number', 'رقم القطعة'), ('name', 'الاسم'), ('category', 'الفئة'),
    ('brand', 'العلامة'), ('oem_number', 'الرقم الأصلي'), ('quality_grade', 'الجودة'),
    ('aliases', 'أسماء أخرى'), ('barcode', 'الباركود'), ('purchase_price', 'سعر الشراء'),
    ('selling_price', 'سعر البيع'), ('quantity', 'الكمية'), ('min_stock_alert', 'حد التنبيه'),
    ('shelf_location', 'الرف'), ('description', 'الوصف'),
]


class ImportFileError(Exception):
    """الملف نفسه غير صالح (صيغة غير مدعومة، بلا عناوين، أكبر من الحد)."""


def read_rows(uploaded_file) -> list:
    """قراءة الملف إلى قائمة صفوف (قوائم قيم)، الصف الأول عناوين."""
    name = (uploaded_file.name or '').lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        except Exception as exc:
            raise ImportFileError('تعذّرت قراءة ملف Excel. احفظه بصيغة .xlsx وحاول مجدداً.') from exc
        sheet = workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif name.endswith('.csv'):
        raw = uploaded_file.read()
        for encoding in ('utf-8-sig', 'cp1256'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportFileError('ترميز ملف CSV غير معروف. احفظه بترميز UTF-8.')
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise ImportFileError('الصيغ المدعومة: Excel (.xlsx) أو CSV.')

    rows = [row for row in rows if any(cell not in (None, '') for cell in row)]
    if not rows:
        raise ImportFileError('الملف فارغ.')
    if len(rows) - 1 > MAX_ROWS:
        raise ImportFileError(f'الحد الأقصى {MAX_ROWS} صف في الملف الواحد.')
    return rows


def _map_headers(header_row) -> dict:
    mapping = {}
    for index, header in enumerate(header_row):
        column = HEADER_LOOKUP.get(normalize_text(header))
        if column and column not in mapping:
            mapping[column] = index
    missing = [label for column, label in (('name', 'الاسم'), ('part_number', 'رقم القطعة'))
               if column not in mapping]
    if missing:
        raise ImportFileError('أعمدة مطلوبة غير موجودة في الصف الأول: ' + '، '.join(missing))
    return mapping


def _text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _decimal(value, label, errors):
    text = _text(value).replace(',', '')
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation:
        errors.append(f'{label} غير رقمي: {text}')
        return None
    if number < 0:
        errors.append(f'{label} لا يمكن أن يكون سالباً.')
        return None
    return number.quantize(Decimal('0.01'))


def _integer(value, label, errors):
    number = _decimal(value, label, errors)
    if number is None:
        return None
    if number != number.to_integral_value():
        errors.append(f'{label} يجب أن يكون عدداً صحيحاً.')
        return None
    return int(number)


def parse_rows(rows) -> list:
    """تحويل الصفوف إلى سجلات مع أخطاء كل صف (رقم الصف كما في Excel)."""
    mapping = _map_headers(rows[0])
    records = []
    seen_numbers = {}
    for offset, row in enumerate(rows[1:], start=2):
        def cell(column):
            index = mapping.get(column)
            return row[index] if index is not None and index < len(row) else None

        errors = []
        record = {
            'row': offset,
            'name': _text(cell('name')),
            'part_number': _text(cell('part_number')),
            'category': _text(cell('category')),
            'brand': _text(cell('brand')),
            'oem_number': _text(cell('oem_number')),
            'aliases': _text(cell('aliases')),
            'barcode': _text(cell('barcode')) or None,
            'shelf_location': _text(cell('shelf_location')),
            'description': _text(cell('description')),
            'purchase_price': _decimal(cell('purchase_price'), 'سعر الشراء', errors),
            'selling_price': _decimal(cell('selling_price'), 'سعر البيع', errors),
            'quantity': _integer(cell('quantity'), 'الكمية', errors),
            'min_stock_alert': _integer(cell('min_stock_alert'), 'حد التنبيه', errors),
        }
        grade_text = _text(cell('quality_grade'))
        record['quality_grade'] = GRADE_LOOKUP.get(normalize_text(grade_text), '') if grade_text else ''
        if grade_text and not record['quality_grade']:
            errors.append(f'درجة الجودة غير معروفة: {grade_text} (أصلي / تجاري / مستعمل).')

        if not record['name']:
            errors.append('الاسم مطلوب.')
        if not record['part_number']:
            errors.append('رقم القطعة مطلوب.')
        elif record['part_number'] in seen_numbers:
            errors.append(f"رقم القطعة مكرر في الملف (الصف {seen_numbers[record['part_number']]}).")
        else:
            seen_numbers[record['part_number']] = offset
        record['errors'] = errors
        records.append(record)
    return records


def plan_import(records, *, update_existing: bool) -> dict:
    """تحديد ما سيُنشأ وما سيُحدَّث، وأخطاء القطع الجديدة الناقصة."""
    numbers = [record['part_number'] for record in records if record['part_number']]
    existing = {part.part_number: part for part in SparePart.objects.filter(part_number__in=numbers)}
    barcodes = {
        record['barcode'] for record in records if record['barcode']
    }
    taken_barcodes = dict(
        SparePart.objects.filter(barcode__in=barcodes).values_list('barcode', 'part_number')
    )

    to_create, to_update, skipped, warnings = [], [], [], []
    for record in records:
        if record['errors']:
            continue
        part = existing.get(record['part_number'])
        owner = taken_barcodes.get(record['barcode'])
        if record['barcode'] and owner and owner != record['part_number']:
            record['errors'].append(f"الباركود مستخدم للقطعة {owner}.")
            continue
        if part is None:
            if record['selling_price'] is None:
                record['errors'].append('سعر البيع مطلوب للقطعة الجديدة.')
            if record['purchase_price'] is None:
                record['errors'].append('سعر الشراء مطلوب للقطعة الجديدة.')
            if not record['category']:
                record['errors'].append('الفئة مطلوبة للقطعة الجديدة.')
            if not record['errors']:
                to_create.append(record)
        elif update_existing:
            to_update.append((part, record))
            if record['quantity'] is not None and record['quantity'] != part.stock_quantity:
                warnings.append(
                    f"الصف {record['row']}: كمية {record['part_number']} لا تُحدَّث من الملف "
                    f"(الحالية {part.stock_quantity}). استخدم الجرد لتصحيحها."
                )
        else:
            skipped.append(record)

    errors = [
        {'row': record['row'], 'part_number': record['part_number'], 'messages': record['errors']}
        for record in records if record['errors']
    ]
    return {
        'to_create': to_create, 'to_update': to_update, 'skipped': skipped,
        'errors': errors, 'warnings': warnings,
    }


def summarize(plan) -> dict:
    return {
        'create_count': len(plan['to_create']),
        'update_count': len(plan['to_update']),
        'skipped_count': len(plan['skipped']),
        'errors': plan['errors'],
        'warnings': plan['warnings'],
        'preview': [
            {'row': record['row'], 'part_number': record['part_number'], 'name': record['name'],
             'action': 'create'}
            for record in plan['to_create'][:50]
        ] + [
            {'row': record['row'], 'part_number': record['part_number'], 'name': record['name'],
             'action': 'update'}
            for _, record in plan['to_update'][:50]
        ],
    }


UPDATABLE_FIELDS = (
    'name', 'brand', 'oem_number', 'quality_grade', 'aliases', 'barcode',
    'shelf_location', 'description', 'selling_price', 'min_stock_alert',
)


def apply_import(plan, *, user) -> None:
    """تطبيق خطة الاستيراد كاملة في معاملة واحدة."""
    categories = {category.name: category for category in Category.objects.all()}

    def category_for(name):
        if name not in categories:
            categories[name] = Category.objects.create(name=name)
        return categories[name]

    with transaction.atomic():
        for record in plan['to_create']:
            part = SparePart.objects.create(
                name=record['name'],
                part_number=record['part_number'],
                category=category_for(record['category']),
                purchase_price=record['purchase_price'],
                selling_price=record['selling_price'],
                min_stock_alert=record['min_stock_alert'] if record['min_stock_alert'] is not None else 5,
                shelf_location=record['shelf_location'],
                brand=record['brand'],
                oem_number=record['oem_number'],
                quality_grade=record['quality_grade'],
                aliases=record['aliases'],
                barcode=record['barcode'],
                description=record['description'] or None,
            )
            if record['quantity']:
                services.set_opening_balance(part, record['quantity'], user=user)

        for part, record in plan['to_update']:
            changed = []
            for field in UPDATABLE_FIELDS:
                value = record[field]
                if value in (None, ''):
                    # الخلية الفارغة لا تمسح قيمة موجودة.
                    continue
                if getattr(part, field) != value:
                    setattr(part, field, value)
                    changed.append(field)
            if record['category'] and part.category.name != record['category']:
                part.category = category_for(record['category'])
                changed.append('category')
            if changed:
                part.save(update_fields=changed)


def export_rows(parts, *, include_costs: bool) -> list:
    columns = [(key, label) for key, label in EXPORT_COLUMNS
               if include_costs or key != 'purchase_price']
    rows = [[label for _, label in columns]]
    for part in parts:
        values = {
            'part_number': part.part_number, 'name': part.name, 'category': part.category.name,
            'brand': part.brand, 'oem_number': part.oem_number,
            'quality_grade': part.get_quality_grade_display() if part.quality_grade else '',
            'aliases': part.aliases, 'barcode': part.barcode or '',
            'purchase_price': float(part.purchase_price), 'selling_price': float(part.selling_price),
            'quantity': part.stock_quantity, 'min_stock_alert': part.min_stock_alert,
            'shelf_location': part.shelf_location, 'description': part.description or '',
        }
        rows.append([values[key] for key, _ in columns])
    return rows


def rows_to_xlsx(rows, title='قطع الغيار') -> bytes:
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.sheet_view.rightToLeft = True
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
