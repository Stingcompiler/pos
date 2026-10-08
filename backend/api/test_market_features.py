"""
اختبارات مزايا السوق السوداني: الدفع المختلط والآجل، رقم الإشعار، التحصيل
وكشف الحساب، المرتجعات، إقفال اليومية، التسعير بسعر الصرف، البحث الموحّد،
الاستيراد والتصدير، الجرد، وإخفاء التكلفة عن الموظف.
"""

import io
import tempfile
import zipfile
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TransactionTestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.token_blacklist.models import BlacklistedToken
from rest_framework_simplejwt.tokens import RefreshToken

from api import services
from api.models import (
    BankAccount,
    Customer,
    DailyClose,
    ExchangeRate,
    Expense,
    Invoice,
    Payment,
    SaleReturn,
    SiteSetting,
    SparePart,
    StockCount,
    StockMovement,
)
from api.search import normalize_text
from api.tests import BaseAPITestCase, User


class MarketTestCase(BaseAPITestCase):
    def setUp(self):
        super().setUp()
        self.bank = BankAccount.objects.create(name='بنكك - بنك الخرطوم', account_number='111')
        self.workshop = Customer.objects.create(
            name='ورشة النور', phone='0911', customer_type=Customer.CustomerType.WORKSHOP,
            credit_limit=Decimal('500.00'),
        )

    def sell(self, client=None, **payload):
        client = client or self.client_for(self.employee)
        body = {'items': [{'spare_part': self.part.pk, 'quantity': 2}], **payload}
        return client.post('/api/invoices/', body, format='json')

    def transfer(self, amount, reference='TRX-1001', **extra):
        return {
            'method': 'bank', 'amount': str(amount), 'bank_account': self.bank.pk,
            'reference_id': reference, 'sender_account_number': '9988', **extra,
        }


class PaymentsAndCreditTests(MarketTestCase):

    def test_mixed_payment_records_each_method(self):
        response = self.sell(payments=[{'method': 'cash', 'amount': '20'}, self.transfer(30)])
        self.assertEqual(response.status_code, 201, response.data)
        invoice = Invoice.objects.get()
        self.assertEqual(invoice.payment_method, Invoice.PaymentMethod.MIXED)
        self.assertEqual(invoice.paid_amount, Decimal('50.00'))
        self.assertEqual(invoice.credit_amount, Decimal('0'))
        self.assertEqual(sorted(invoice.payments.values_list('method', flat=True)), ['bank', 'cash'])
        self.assertEqual(invoice.reference_id, 'TRX-1001')
        self.assertEqual(len(response.data['payments']), 2)

    def test_transfer_reference_cannot_be_reused(self):
        self.assertEqual(self.sell(payments=[self.transfer(50, 'TRX 55-01')]).status_code, 201)
        # الصيغة نفسها بمسافات أو شرطات مختلفة تُعدّ الإشعار نفسه.
        response = self.sell(payments=[self.transfer(50, 'trx5501')])
        self.assertEqual(response.status_code, 400)
        self.assertIn('payment', response.data)
        self.assertIn('مستخدم مسبقاً', response.data['payment'])
        self.assertEqual(Invoice.objects.count(), 1)
        self.part.refresh_from_db()
        self.assertEqual(self.part.stock_quantity, 48)

    def test_legacy_bank_payload_still_works_and_is_checked(self):
        legacy = {
            'payment_method': 'bank', 'bank_name': 'بنكك', 'reference_id': 'OLD-1',
            'sender_account_number': '77',
        }
        self.assertEqual(self.sell(**legacy).status_code, 201)
        self.assertEqual(Payment.objects.get().reference_key, 'OLD1')
        self.assertEqual(self.sell(**legacy).status_code, 400)

    def test_overpayment_is_rejected(self):
        response = self.sell(payments=[{'method': 'cash', 'amount': '60'}])
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Invoice.objects.count(), 0)

    def test_credit_sale_requires_customer_with_limit(self):
        self.assertEqual(self.sell(payment_method='credit').status_code, 400)
        retail = Customer.objects.create(name='زبون', phone='0900')
        response = self.sell(payment_method='credit', customer=retail.pk)
        self.assertEqual(response.status_code, 400)
        self.assertIn('غير مسموح', response.data['payment'])

    def test_partial_credit_sale_within_limit(self):
        # خصم الورشة 10% من الإعدادات: 2 × 25 × 0.9 = 45
        site = SiteSetting.load()
        site.workshop_discount_percent = Decimal('10')
        site.save()
        response = self.sell(customer=self.workshop.pk, payments=[{'method': 'cash', 'amount': '15'}])
        self.assertEqual(response.status_code, 201, response.data)
        invoice = Invoice.objects.get()
        self.assertEqual(invoice.total_amount, Decimal('45.00'))
        self.assertEqual(invoice.credit_amount, Decimal('30.00'))
        self.assertEqual(invoice.payment_method, Invoice.PaymentMethod.CREDIT)
        self.assertEqual(services.customer_balance(self.workshop), Decimal('30.00'))

    def test_credit_limit_cannot_be_exceeded(self):
        self.workshop.credit_limit = Decimal('40.00')
        self.workshop.save()
        response = self.sell(customer=self.workshop.pk, payment_method='credit')
        self.assertEqual(response.status_code, 400)
        self.assertIn('حد الائتمان', response.data['payment'])
        self.part.refresh_from_db()
        self.assertEqual(self.part.stock_quantity, 50)

    def test_collection_statement_and_balance(self):
        self.sell(customer=self.workshop.pk, payment_method='credit')
        client = self.client_for(self.employee)
        over = client.post(f'/api/customers/{self.workshop.pk}/payments/',
                           {'method': 'cash', 'amount': '60'}, format='json')
        self.assertEqual(over.status_code, 400)
        paid = client.post(f'/api/customers/{self.workshop.pk}/payments/',
                           {'method': 'cash', 'amount': '20'}, format='json')
        self.assertEqual(paid.status_code, 201, paid.data)

        statement = self.client_for(self.manager).get(f'/api/customers/{self.workshop.pk}/statement/')
        self.assertEqual([entry['type'] for entry in statement.data['entries']], ['invoice', 'collection'])
        self.assertEqual(statement.data['entries'][-1]['balance'], '30.00')
        listing = self.client_for(self.manager).get('/api/customers/', {'has_balance': 1})
        self.assertEqual(listing.data['results'][0]['balance'], '30.00')

    def test_employee_adds_customer_without_credit_terms(self):
        response = self.client_for(self.employee).post('/api/customers/', {
            'name': 'عميل جديد', 'phone': '0123', 'credit_limit': '99999', 'discount_percent': '50',
        }, format='json')
        self.assertEqual(response.status_code, 201)
        customer = Customer.objects.get(name='عميل جديد')
        self.assertEqual(customer.credit_limit, Decimal('0'))
        self.assertIsNone(customer.discount_percent)

    def test_customer_with_invoices_cannot_be_deleted(self):
        self.sell(customer=self.workshop.pk, payment_method='credit')
        response = self.client_for(self.manager).delete(f'/api/customers/{self.workshop.pk}/')
        self.assertEqual(response.status_code, 400)


class SaleReturnTests(MarketTestCase):

    def setUp(self):
        super().setUp()
        self.invoice = services.create_invoice(
            cashier=self.employee, items=[{'spare_part': self.part, 'quantity': 4}],
        )
        self.item = self.invoice.items.get()

    def return_items(self, quantity, refund_method='cash', user=None, **extra):
        client = self.client_for(user or self.manager)
        return client.post(f'/api/invoices/{self.invoice.pk}/returns/', {
            'items': [{'invoice_item': self.item.pk, 'quantity': quantity}],
            'refund_method': refund_method, **extra,
        }, format='json')

    def test_partial_return_restocks_and_refunds(self):
        response = self.return_items(1, reason='مقاس خطأ')
        self.assertEqual(response.status_code, 201, response.data)
        self.part.refresh_from_db()
        self.assertEqual(self.part.stock_quantity, 47)
        movement = StockMovement.objects.get(reason=StockMovement.Reason.RETURN)
        self.assertEqual(movement.unit_cost, Decimal('10.00'))
        refund = Payment.objects.get(kind=Payment.Kind.REFUND)
        self.assertEqual(refund.amount, Decimal('25.00'))

    def test_cannot_return_more_than_sold_minus_returned(self):
        self.assertEqual(self.return_items(3).status_code, 201)
        response = self.return_items(2)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(SaleReturn.objects.count(), 1)

    def test_employee_cannot_create_return(self):
        self.assertEqual(self.return_items(1, user=self.employee).status_code, 403)

    def test_credit_invoice_refund_goes_to_account(self):
        invoice = services.create_invoice(
            cashier=self.employee, customer=self.workshop, payment_method='credit',
            items=[{'spare_part': self.part, 'quantity': 2}],
        )
        item = invoice.items.get()
        with self.assertRaises(services.PaymentError):
            services.create_sale_return(
                invoice=invoice, items=[{'invoice_item': item, 'quantity': 1}],
                refund_method='cash', user=self.manager,
            )
        services.create_sale_return(
            invoice=invoice, items=[{'invoice_item': item, 'quantity': 1}],
            refund_method='account', user=self.manager,
        )
        self.assertEqual(services.customer_balance(self.workshop), Decimal('25.00'))

    def test_reports_are_net_of_returns(self):
        self.return_items(1)
        overall = self.client_for(self.manager).get('/api/reports/sales/').data['overall']
        self.assertEqual(overall['gross_revenue'], 100.0)
        self.assertEqual(overall['returns_total'], 25.0)
        self.assertEqual(overall['total_revenue'], 75.0)
        # ربح البند 15 للوحدة: 4 مباعة - 1 مرتجعة = 45
        self.assertEqual(overall['total_profit'], 45.0)
        detail = self.client_for(self.manager).get(f'/api/invoices/{self.invoice.pk}/').data
        self.assertEqual(detail['items'][0]['returned_quantity'], 1)
        self.assertEqual(len(detail['returns']), 1)


class DailyCloseTests(MarketTestCase):

    def test_summary_and_close(self):
        self.sell(payments=[{'method': 'cash', 'amount': '30'}, self.transfer(20)])
        Expense.objects.create(date=timezone.localdate(), category='كهرباء', amount=Decimal('5'),
                               created_by=self.manager)
        client = self.client_for(self.manager)
        summary = client.get('/api/daily-summary/', {'opening_cash': '100'}).data
        self.assertEqual(Decimal(summary['expected_cash']), Decimal('125.00'))
        self.assertEqual(summary['banks'][0]['in'], Decimal('20.00'))
        self.assertEqual(summary['banks'][0]['unverified'], 1)

        closed = client.post('/api/daily-closes/', {'counted_cash': '120', 'opening_cash': '100'},
                             format='json')
        self.assertEqual(closed.status_code, 201, closed.data)
        self.assertEqual(Decimal(closed.data['difference']), Decimal('-5.00'))
        again = client.post('/api/daily-closes/', {'counted_cash': '120'}, format='json')
        self.assertEqual(again.status_code, 400)
        # الإقفال التالي يبدأ بنقد اليوم المعدود.
        tomorrow = timezone.localdate() + timedelta(days=1)
        self.assertEqual(services.default_opening_cash(tomorrow), Decimal('120.00'))

    def test_closed_day_rejects_expenses(self):
        DailyClose.objects.create(date=timezone.localdate(), expected_cash=0, counted_cash=0,
                                  difference=0, closed_by=self.manager)
        response = self.client_for(self.manager).post('/api/expenses/', {
            'date': timezone.localdate().isoformat(), 'category': 'ترحيل', 'amount': '10',
        }, format='json')
        self.assertEqual(response.status_code, 400)

    def test_transfer_verification_is_privileged(self):
        self.sell(payments=[self.transfer(50)])
        payment = Payment.objects.get()
        self.assertEqual(self.client_for(self.employee).post(f'/api/payments/{payment.pk}/verify/').status_code, 403)
        response = self.client_for(self.supervisor).post(f'/api/payments/{payment.pk}/verify/')
        self.assertEqual(response.status_code, 200)
        self.assertIsNotNone(response.data['verified_at'])


class PricingTests(MarketTestCase):

    def test_foreign_supply_sets_cost_and_reprice(self):
        services.record_exchange_rate(currency='USD', rate=Decimal('2000'), user=self.manager)
        services.create_supply_deal(
            supplier=self.supplier, spare_part=self.part, quantity_added=10,
            currency='USD', foreign_unit_cost=Decimal('5'), user=self.manager,
        )
        self.part.refresh_from_db()
        self.assertEqual(self.part.cost_currency, 'USD')
        self.assertEqual(self.part.foreign_cost, Decimal('5.00'))

        site = SiteSetting.load()
        site.default_markup_percent = Decimal('20')
        site.price_rounding = 50
        site.save()
        services.record_exchange_rate(currency='USD', rate=Decimal('2100'), user=self.manager)

        client = self.client_for(self.manager)
        preview = client.get('/api/pricing/').data
        # 5 × 2100 × 1.2 = 12600 (مضاعف 50)
        self.assertEqual(preview['changes'][0]['new_price'], '12600')
        self.part.refresh_from_db()
        self.assertEqual(self.part.selling_price, Decimal('25.00'))

        applied = client.post('/api/pricing/', {}, format='json')
        self.assertEqual(applied.data['count'], 1)
        self.part.refresh_from_db()
        self.assertEqual(self.part.selling_price, Decimal('12600.00'))
        self.assertIsNotNone(self.part.price_updated_at)

    def test_unsupported_currency_and_missing_rate(self):
        with self.assertRaises(services.InventoryError):
            services.record_exchange_rate(currency='XYZ', rate=1, user=self.manager)
        with self.assertRaises(services.InventoryError):
            services.create_supply_deal(
                supplier=self.supplier, spare_part=self.part, quantity_added=1,
                currency='AED', foreign_unit_cost=Decimal('3'), user=self.manager,
            )
        self.assertEqual(ExchangeRate.objects.count(), 0)

    def test_employee_cannot_see_costs(self):
        client = self.client_for(self.employee)
        part = client.get(f'/api/spare-parts/{self.part.pk}/').data
        for field in ('purchase_price', 'foreign_cost', 'markup_percent'):
            self.assertNotIn(field, part)
        search = client.get('/api/spare-parts/search-pos/', {'q': 'فلتر'}).data
        self.assertNotIn('purchase_price', search[0])
        services.create_invoice(cashier=self.employee, items=[{'spare_part': self.part, 'quantity': 1}])
        movement = client.get('/api/stock-movements/').data['results'][0]
        self.assertNotIn('unit_cost', movement)
        self.assertIn('unit_cost', self.client_for(self.manager).get('/api/stock-movements/').data['results'][0])
        self.assertNotIn('default_markup_percent', client.get('/api/admin/settings/').data)


class SearchAndCatalogTests(MarketTestCase):

    def test_normalization(self):
        self.assertEqual(normalize_text('إطَارٌ أمامي'), 'اطار امامي')
        self.assertEqual(normalize_text('بطارية ٧٠ أمبير'), 'بطاريه 70 امبير')

    def test_search_by_alias_spelling_and_compact_number(self):
        part = SparePart.objects.create(
            name='مضخة مياه', part_number='16100-39465', category=self.category,
            purchase_price=Decimal('1'), selling_price=Decimal('2'), stock_quantity=3,
            aliases='طرمبة ماء\nواتر بمب', brand='Aisin', quality_grade='original',
        )
        part.compatible_cars.add(self.car_model)
        client = self.client_for(self.employee)
        for query in ('طرمبه', 'مضخه', '16100 39465', '1610039465', 'Corolla', 'aisin اصلي'):
            ids = [row['id'] for row in client.get('/api/spare-parts/', {'search': query}).data['results']]
            self.assertIn(part.pk, ids, query)

    def test_barcode_lookup_and_pos_exact_match(self):
        self.part.barcode = '6291041500213'
        self.part.save()
        client = self.client_for(self.employee)
        self.assertEqual(client.get('/api/spare-parts/lookup/', {'code': '6291041500213'}).data['id'], self.part.pk)
        self.assertEqual(client.get('/api/spare-parts/lookup/', {'code': 'of 100'}).data['id'], self.part.pk)
        self.assertEqual(client.get('/api/spare-parts/lookup/', {'code': 'none'}).status_code, 404)

    def test_opening_quantity_creates_movement(self):
        response = self.client_for(self.manager).post('/api/spare-parts/', {
            'name': 'بوجي', 'part_number': 'SP-9', 'category': self.category.pk,
            'purchase_price': '3', 'selling_price': '5', 'opening_quantity': 12,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['stock_quantity'], 12)
        movement = StockMovement.objects.get(reason=StockMovement.Reason.OPENING)
        self.assertEqual(movement.change, 12)
        self.assertEqual(movement.unit_cost, Decimal('3.00'))


class ImportExportTests(MarketTestCase):

    def workbook(self, rows):
        book = Workbook()
        for row in rows:
            book.active.append(row)
        output = io.BytesIO()
        book.save(output)
        return SimpleUploadedFile('parts.xlsx', output.getvalue())

    def upload(self, rows, **params):
        return self.client_for(self.manager).post(
            '/api/spare-parts/import/', {'file': self.workbook(rows), **params}, format='multipart',
        )

    def test_preview_then_apply(self):
        rows = [
            ['رقم القطعة', 'الاسم', 'الفئة', 'سعر الشراء', 'سعر البيع', 'الكمية', 'الجودة', 'أسماء أخرى'],
            ['BK-1', 'فحمة فرامل', 'فرامل', 10, 18, 6, 'تجاري', 'تيل فرامل'],
            ['OF-100', 'فلتر زيت محدّث', '', '', 30, 99, '', ''],
        ]
        preview = self.upload(rows)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual((preview.data['create_count'], preview.data['update_count']), (1, 0))
        self.assertFalse(SparePart.objects.filter(part_number='BK-1').exists())

        applied = self.upload(rows, dry_run='0', update_existing='1')
        self.assertTrue(applied.data['applied'], applied.data)
        part = SparePart.objects.get(part_number='BK-1')
        self.assertEqual((part.stock_quantity, part.quality_grade, part.category.name), (6, 'commercial', 'فرامل'))
        self.part.refresh_from_db()
        # التحديث يغيّر الاسم وسعر البيع، ولا يلمس الرصيد أو التكلفة.
        self.assertEqual((self.part.name, self.part.selling_price, self.part.stock_quantity),
                         ('فلتر زيت محدّث', Decimal('30.00'), 50))
        self.assertEqual(len(applied.data['warnings']), 1)

    def test_errors_block_apply(self):
        rows = [['رقم القطعة', 'الاسم', 'سعر البيع'], ['X-1', 'قطعة', 'abc'], ['', 'بلا رقم', 5]]
        response = self.upload(rows, dry_run='0')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(len(response.data['errors']), 2)
        self.assertFalse(SparePart.objects.filter(part_number='X-1').exists())

    def test_export_hides_cost_from_employee(self):
        manager_file = self.client_for(self.manager).get('/api/spare-parts/export/')
        employee_file = self.client_for(self.employee).get('/api/spare-parts/export/')
        manager_header = [c.value for c in load_workbook(io.BytesIO(manager_file.content)).active[1]]
        employee_header = [c.value for c in load_workbook(io.BytesIO(employee_file.content)).active[1]]
        self.assertIn('سعر الشراء', manager_header)
        self.assertNotIn('سعر الشراء', employee_header)


class StockCountTests(MarketTestCase):

    def test_count_and_apply(self):
        self.part.barcode = '111222'
        self.part.save()
        client = self.client_for(self.supervisor)
        count = client.post('/api/stock-counts/', {'title': 'جرد الرف A'}, format='json').data
        for _ in range(3):
            client.post(f"/api/stock-counts/{count['id']}/lines/",
                        {'code': '111222', 'counted_quantity': 1, 'mode': 'add'}, format='json')
        detail = client.get(f"/api/stock-counts/{count['id']}/").data
        self.assertEqual(detail['lines'][0]['counted_quantity'], 3)

        applied = client.post(f"/api/stock-counts/{count['id']}/apply/")
        self.assertEqual(applied.status_code, 200, applied.data)
        self.part.refresh_from_db()
        self.assertEqual(self.part.stock_quantity, 3)
        movement = StockMovement.objects.get(reason=StockMovement.Reason.STOCK_COUNT)
        self.assertEqual(movement.change, -47)
        self.assertEqual(StockCount.objects.get().lines.get().system_quantity, 50)
        self.assertEqual(client.post(f"/api/stock-counts/{count['id']}/apply/").status_code, 400)

    def test_employee_cannot_run_counts(self):
        self.assertEqual(self.client_for(self.employee).post('/api/stock-counts/', {}, format='json').status_code, 403)


class AnonymousAccessTests(MarketTestCase):

    def test_receipt_settings_for_all_staff(self):
        response = self.client_for(self.employee).get('/api/receipt-settings/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('receipt_paper', response.data)
        self.assertEqual(APIClient().get('/api/receipt-settings/').status_code, 401)


class BackupRestoreTests(TransactionTestCase):
    """نسخة → مسح → استعادة تعيد البيانات نفسها، ونسخة تالفة تُرفض بلا كتابة."""

    def test_round_trip(self):
        from django.contrib.auth import get_user_model
        from api.models import Category

        user = get_user_model().objects.create_user(username='owner', password='StrongPass123', role='manager')
        category = Category.objects.create(name='بطاريات')
        part = SparePart.objects.create(
            name='بطارية 70', part_number='BAT-70', category=category,
            purchase_price=Decimal('50'), selling_price=Decimal('80'), stock_quantity=4,
        )
        services.create_invoice(cashier=user, items=[{'spare_part': part, 'quantity': 1}])
        # جلسة خرج صاحبها: تبقى مُلغاة بعد الاستعادة.
        logged_out = RefreshToken.for_user(user)
        logged_out.blacklist()

        with tempfile.TemporaryDirectory() as folder, override_settings(MEDIA_ROOT=folder + '/media'):
            call_command('backup_data', output_dir=folder, stdout=io.StringIO())
            archive = next(Path(folder).glob('backup-*.zip'))

            with self.assertRaises(CommandError):
                call_command('restore_data', str(archive), stdout=io.StringIO())

            call_command('restore_data', str(archive), force=True, stdout=io.StringIO())
            self.assertEqual(SparePart.objects.get().stock_quantity, 3)
            self.assertEqual(Invoice.objects.get().payments.count(), 1)
            self.assertTrue(BlacklistedToken.objects.filter(token__jti=logged_out['jti']).exists())

            corrupted = Path(folder) / 'backup-corrupted.zip'
            data = zipfile.ZipFile(archive).read('data.json').replace(b'BAT-70', b'BAT-99')
            with zipfile.ZipFile(archive) as source, zipfile.ZipFile(corrupted, 'w') as target:
                target.writestr('manifest.json', source.read('manifest.json'))
                target.writestr('data.json', data)
            with self.assertRaises(CommandError):
                call_command('restore_data', str(corrupted), force=True, stdout=io.StringIO())
            self.assertTrue(SparePart.objects.filter(part_number='BAT-70').exists())


class PrivateProofImageTests(MarketTestCase):
    """صورة الإشعار خارج /media/ العامة، وتُقرأ عبر نقطة نهاية بصلاحيات."""

    def test_proof_served_only_to_authorized_users(self):
        import io as _io
        from PIL import Image

        with tempfile.TemporaryDirectory() as folder, override_settings(PRIVATE_MEDIA_ROOT=folder):
            self.sell(payments=[self.transfer(50)])
            payment = Payment.objects.get()
            buffer = _io.BytesIO()
            Image.new('RGB', (20, 20), 'white').save(buffer, 'PNG')
            upload = SimpleUploadedFile('proof.png', buffer.getvalue(), content_type='image/png')
            owner = self.client_for(self.employee)
            self.assertEqual(owner.post(f'/api/payments/{payment.pk}/proof/', {'proof_image': upload},
                                        format='multipart').status_code, 200)

            payment.refresh_from_db()
            self.assertTrue(Path(folder, payment.proof_image.name).exists())
            url = owner.get(f'/api/invoices/{payment.invoice_id}/').data['payments'][0]['proof_image']
            self.assertEqual(url, f'/api/payments/{payment.pk}/proof-image/')
            self.assertEqual(owner.get(url).status_code, 200)
            self.assertEqual(self.client_for(self.manager).get(url).status_code, 200)

            other = User.objects.create_user(username='emp2', password='StrongPass123', role='employee')
            self.assertEqual(self.client_for(other).get(url).status_code, 404)
            self.assertEqual(APIClient().get(url).status_code, 401)


class PriceChangeDuringCheckoutTests(MarketTestCase):
    """سعر تغيّر بين البحث والبيع: يُرفض البيع بدل دين صامت أو رفض غامض."""

    def test_stale_total_is_rejected_with_409(self):
        SparePart.objects.filter(pk=self.part.pk).update(selling_price=Decimal('30.00'))
        response = self.sell(expected_total='50.00', payments=[{'method': 'cash', 'amount': '50'}])
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.data['total'], '60.00')
        self.assertEqual(Invoice.objects.count(), 0)

    def test_matching_total_sells(self):
        response = self.sell(expected_total='50.00', payments=[{'method': 'cash', 'amount': '50'}])
        self.assertEqual(response.status_code, 201, response.data)

    def test_two_transfers_on_one_invoice(self):
        response = self.sell(payments=[self.transfer(30, 'TRX-A'), self.transfer(20, 'TRX-B')])
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Invoice.objects.get().payment_method, Invoice.PaymentMethod.MIXED)

    def test_closed_day_blocks_sales(self):
        DailyClose.objects.create(date=timezone.localdate(), expected_cash=0, counted_cash=0,
                                  difference=0, closed_by=self.manager)
        response = self.sell(payments=[{'method': 'cash', 'amount': '50'}])
        self.assertEqual(response.status_code, 400)
        self.assertIn('مُقفلة', response.data['payment'])


class StockCountSnapshotTests(MarketTestCase):

    def test_sale_between_count_and_apply_is_kept(self):
        client = self.client_for(self.supervisor)
        count = client.post('/api/stock-counts/', {'title': 'رف A'}, format='json').data
        # العدّ وجد 48 والنظام 50 (قطعتان مفقودتان)، ثم بيعت 5 قبل التطبيق.
        client.post(f"/api/stock-counts/{count['id']}/lines/",
                    {'spare_part': self.part.pk, 'counted_quantity': 48, 'mode': 'set'}, format='json')
        services.create_invoice(cashier=self.employee, items=[{'spare_part': self.part, 'quantity': 5}])

        self.assertEqual(client.post(f"/api/stock-counts/{count['id']}/apply/").status_code, 200)
        self.part.refresh_from_db()
        self.assertEqual(self.part.stock_quantity, 43)
        self.assertEqual(StockMovement.objects.get(reason=StockMovement.Reason.STOCK_COUNT).change, -2)
